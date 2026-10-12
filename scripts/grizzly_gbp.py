#!/usr/bin/env python
from __future__ import annotations

import argparse
import base64
import csv
import dataclasses
import datetime as dt
import hashlib
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

PLUGIN_ROOT = Path(__file__).resolve().parents[1]
CONFIG_EXAMPLE = PLUGIN_ROOT / "config.example.json"
CONFIG_LOCAL = PLUGIN_ROOT / "config.local.json"
DEFAULT_HEADERS = [
    "Date",
    "PostType",
    "Topic",
    "AssetSource",
    "AssetIdOrDescription",
    "CTA",
    "Status",
    "CaptionDraft",
    "ImageLink",
    "Posted",
    "PostedAt",
    "GBPPostUrl",
    "Notes",
]
ELECTRICAL_EXTENSIONS = {".jpg", ".jpeg", ".png", ".heic", ".heif"}


@dataclasses.dataclass
class PhotoScore:
    score: int
    electrical: bool
    usable: bool
    reasons: list[str]
    topic_hint: str


def now_stamp() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def load_config() -> dict[str, Any]:
    path = CONFIG_LOCAL if CONFIG_LOCAL.exists() else CONFIG_EXAMPLE
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def save_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=True)
        handle.write("\n")
    tmp.replace(path)


def load_state(config: dict[str, Any]) -> dict[str, Any]:
    path = Path(config["state_path"])
    if not path.exists():
        return {"processed_photos": {}, "copied_candidates": [], "posted_rows": {}}
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def log(config: dict[str, Any], message: str) -> None:
    folder = Path(config["log_folder"])
    folder.mkdir(parents=True, exist_ok=True)
    line = f"[{now_stamp()}] {message}"
    print(line)
    with (folder / f"{dt.date.today().isoformat()}.log").open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def init_config(_: argparse.Namespace) -> None:
    if CONFIG_LOCAL.exists():
        print(f"Already exists: {CONFIG_LOCAL}")
        return
    shutil.copy2(CONFIG_EXAMPLE, CONFIG_LOCAL)
    print(f"Created {CONFIG_LOCAL}")


def image_files_in(folder: Path, recursive: bool) -> list[Path]:
    pattern = "**/*" if recursive else "*"
    return [
        p
        for p in folder.glob(pattern)
        if p.is_file() and p.suffix.lower() in ELECTRICAL_EXTENSIONS
    ]


def file_identity(path: Path) -> str:
    stat = path.stat()
    raw = f"{path.resolve()}|{stat.st_size}|{int(stat.st_mtime)}"
    return hashlib.sha256(raw.encode("utf-8", errors="ignore")).hexdigest()


def short_hash(path: Path) -> str:
    stat = path.stat()
    raw = f"{path.name}|{stat.st_size}|{int(stat.st_mtime)}"
    return hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:10]


def slugify(value: str, fallback: str = "photo") -> str:
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return value[:80] or fallback


def parse_date_from_name(name: str) -> dt.date | None:
    match = re.search(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})", name)
    if not match:
        return None
    try:
        return dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def guess_topic_from_name(path: Path, keywords: list[str]) -> str:
    stem = path.stem.lower()
    hits = [word for word in keywords if word.lower() in stem]
    if hits:
        return " / ".join(dict.fromkeys(hits[:4]))
    cleaned = re.sub(r"20\d{2}[-_]?\d{2}[-_]?\d{2}|img|ios|photo|copy", " ", stem)
    cleaned = re.sub(r"[^a-z0-9]+", " ", cleaned).strip()
    return cleaned[:60] or "recent electrical work"


def heuristic_score(path: Path, config: dict[str, Any]) -> PhotoScore:
    reasons: list[str] = []
    score = 0
    name = path.name.lower()
    keywords = config.get("electrical_keywords", [])
    hits = [word for word in keywords if word.lower() in name]
    if hits:
        score += 50
        reasons.append(f"name keywords: {', '.join(hits[:5])}")
    if parse_date_from_name(path.name):
        score += 10
        reasons.append("date-like filename")
    try:
        size = path.stat().st_size
    except OSError:
        size = 0
    if size >= 250_000:
        score += 15
        reasons.append("reasonable image size")
    elif size:
        reasons.append("small image")
    lowered = path.name.lower()
    bad_words = ["screenshot", "screenrecording", "receipt", "refund", "qr", "logo"]
    if any(word in lowered for word in bad_words):
        score -= 25
        reasons.append("likely non-job media")
    topic_hint = guess_topic_from_name(path, keywords)
    electrical = score >= 50
    usable = score >= int(config.get("minimum_photo_score", 65))
    return PhotoScore(max(0, min(100, score)), electrical, usable, reasons, topic_hint)


def openai_score(path: Path, config: dict[str, Any]) -> PhotoScore | None:
    api_key = config.get("openai_api_key") or os.environ.get("OPENAI_API_KEY")
    if not api_key or path.suffix.lower() not in {".jpg", ".jpeg", ".png"}:
        return None
    try:
        data = base64.b64encode(path.read_bytes()).decode("ascii")
    except OSError:
        return None
    mime = mimetypes.guess_type(path.name)[0] or "image/jpeg"
    prompt = (
        "Score this image for a Grizzly Electrical Solutions Google Business Profile post. "
        "Return strict JSON with score 0-100, electrical boolean, usable boolean, "
        "topic_hint string, and reasons array. Penalize faces, customer personal info, "
        "unsafe work, blurry/poor framing, memes, screenshots, logos, receipts, and non-electrical content."
    )
    payload = {
        "model": config.get("openai_model", "gpt-5.4-mini"),
        "input": [
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": prompt},
                    {"type": "input_image", "image_url": f"data:{mime};base64,{data}"},
                ],
            }
        ],
        "text": {"format": {"type": "json_object"}},
    }
    request = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            result = json.loads(response.read().decode("utf-8"))
    except Exception:
        return None
    text = result.get("output_text")
    if not text:
        chunks: list[str] = []
        for item in result.get("output", []):
            for content in item.get("content", []):
                if "text" in content:
                    chunks.append(content["text"])
        text = "".join(chunks)
    try:
        parsed = json.loads(text)
        return PhotoScore(
            int(parsed.get("score", 0)),
            bool(parsed.get("electrical", False)),
            bool(parsed.get("usable", False)),
            list(parsed.get("reasons", []))[:8],
            str(parsed.get("topic_hint") or "recent electrical work"),
        )
    except Exception:
        return None


def iter_raw_candidates(config: dict[str, Any]) -> list[Path]:
    raw = Path(config["raw_photo_folder"])
    recursive = bool(config.get("raw_scan_recursive", False))
    excluded = {name.lower() for name in config.get("excluded_folder_names", [])}
    recent_days = int(config.get("scan_recent_days", 60))
    cutoff = time.time() - recent_days * 86400
    files: list[Path] = []
    for path in image_files_in(raw, recursive=recursive):
        rel_parts = [part.lower() for part in path.relative_to(raw).parts[:-1]]
        if any(part in excluded for part in rel_parts):
            continue
        if path.parent != raw and not recursive:
            continue
        try:
            if path.stat().st_mtime < cutoff:
                continue
        except OSError:
            continue
        files.append(path)
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[: int(config.get("max_scan_files", 75))]


def copy_candidate(path: Path, score: PhotoScore, config: dict[str, Any], dry_run: bool) -> Path:
    curated = Path(config["curated_photo_folder"])
    source_date = parse_date_from_name(path.name) or dt.date.today()
    name = f"{source_date.isoformat()}-{slugify(score.topic_hint)}-{short_hash(path)}{path.suffix.lower()}"
    target = curated / name
    counter = 2
    while target.exists():
        target = curated / f"{source_date.isoformat()}-{slugify(score.topic_hint)}-{short_hash(path)}-{counter}{path.suffix.lower()}"
        counter += 1
    if not dry_run:
        curated.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)
    return target


def scan_photos(args: argparse.Namespace) -> None:
    config = load_config()
    if args.recent_days:
        config["scan_recent_days"] = args.recent_days
    state = load_state(config)
    copied: list[dict[str, Any]] = []
    for path in iter_raw_candidates(config):
        try:
            identity = file_identity(path)
        except OSError as exc:
            message = f"SKIP unreadable {path}: {exc}"
            print(message) if args.dry_run else log(config, message)
            continue
        if identity in state.get("processed_photos", {}):
            continue
        if args.dry_run:
            score = heuristic_score(path, config)
        else:
            score = openai_score(path, config) or heuristic_score(path, config)
        record = {
            "path": str(path),
            "processed_at": now_stamp(),
            "score": score.score,
            "electrical": score.electrical,
            "usable": score.usable,
            "reasons": score.reasons,
            "topic_hint": score.topic_hint,
            "copied_to": None,
        }
        if score.usable:
            target = copy_candidate(path, score, config, args.dry_run)
            record["copied_to"] = str(target)
            copied.append(record)
            message = f"{'DRY ' if args.dry_run else ''}COPY candidate {path.name} -> {target.name} ({score.score})"
            print(message) if args.dry_run else log(config, message)
        else:
            message = f"SKIP candidate {path.name} score={score.score} reasons={'; '.join(score.reasons)}"
            print(message) if args.dry_run else log(config, message)
        if not args.dry_run:
            state.setdefault("processed_photos", {})[identity] = record
    if not args.dry_run:
        state.setdefault("copied_candidates", []).extend(copied)
        save_json(Path(config["state_path"]), state)
    print(f"Candidates copied: {len(copied)}")


def open_posts_sheet(config: dict[str, Any]):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(
            "openpyxl is required for workbook commands. Use the Codex bundled Python runtime "
            "or install openpyxl for this Python."
        ) from exc
    workbook_path = Path(config["workbook_path"])
    wb = load_workbook(workbook_path)
    if "Posts" not in wb.sheetnames:
        raise SystemExit("Workbook must contain a Posts sheet.")
    ws = wb["Posts"]
    headers = [cell.value for cell in ws[1]]
    missing = [h for h in DEFAULT_HEADERS if h not in headers]
    if missing:
        raise SystemExit(f"Workbook missing headers: {missing}")
    cols = {name: headers.index(name) + 1 for name in DEFAULT_HEADERS}
    return wb, ws, cols


def row_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        try:
            return dt.date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def find_row_by_date(ws, cols: dict[str, int], target: dt.date) -> int | None:
    for row in range(2, ws.max_row + 1):
        if row_date(ws.cell(row, cols["Date"]).value) == target:
            return row
    return None


def curated_image_for_date(config: dict[str, Any], target: dt.date) -> Path | None:
    curated = Path(config["curated_photo_folder"])
    prefix = target.isoformat()
    dated = sorted(p for p in image_files_in(curated, recursive=False) if p.name.startswith(prefix))
    if dated:
        return dated[0]
    all_images = sorted(image_files_in(curated, recursive=False), key=lambda p: p.stat().st_mtime, reverse=True)
    return all_images[0] if all_images else None


def caption_for(topic: str, cta: str, image: Path | None) -> str:
    image_note = f" This project photo shows {image.stem.replace('-', ' ')}." if image else ""
    return (
        f"{topic}. Reliable electrical work starts with clean details, clear communication, "
        f"and a licensed pro who treats your home or business with care.{image_note}\n\n"
        f"{cta} with Grizzly Electrical Solutions."
    )


def prepare_today(args: argparse.Namespace) -> None:
    config = load_config()
    target = dt.date.fromisoformat(args.date) if args.date else dt.date.today()
    wb, ws, cols = open_posts_sheet(config)
    row = find_row_by_date(ws, cols, target)
    if not row:
        raise SystemExit(f"No Posts row found for {target.isoformat()}.")
    image = curated_image_for_date(config, target)
    topic = str(ws.cell(row, cols["Topic"]).value or "Electrical project update")
    cta = str(ws.cell(row, cols["CTA"]).value or "Contact us today")
    caption = caption_for(topic, cta, image)
    notes = []
    if image:
        notes.append(f"Prepared with image {image.name}")
    else:
        notes.append("No curated image found")
    if args.dry_run:
        print(f"Would prepare row {row} for {target}: {caption[:120]}...")
        return
    ws.cell(row, cols["CaptionDraft"]).value = caption
    ws.cell(row, cols["Status"]).value = config.get("prepared_status", "Needs approval")
    if image:
        ws.cell(row, cols["AssetSource"]).value = "Curated OneDrive"
        ws.cell(row, cols["AssetIdOrDescription"]).value = str(image)
    ws.cell(row, cols["Notes"]).value = "; ".join(notes)
    wb.save(config["workbook_path"])
    log(config, f"Prepared workbook row {row} for {target.isoformat()}")


def month_dates(month: str, days: int | None) -> list[dt.date]:
    year, mon = [int(part) for part in month.split("-")]
    first = dt.date(year, mon, 1)
    if days:
        count = days
    else:
        next_month = dt.date(year + (mon == 12), 1 if mon == 12 else mon + 1, 1)
        count = (next_month - first).days
    return [first + dt.timedelta(days=i) for i in range(count)]


def create_month_schedule(args: argparse.Namespace) -> None:
    config = load_config()
    dates = month_dates(args.month, args.days)
    wb, ws, cols = open_posts_sheet(config)
    existing_rows = []
    existing_by_date: dict[dt.date, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        values = {h: ws.cell(row, cols[h]).value for h in DEFAULT_HEADERS}
        existing_rows.append(values)
        existing_date = row_date(values.get("Date"))
        if existing_date:
            existing_by_date[existing_date] = values
    templates = existing_rows or [{h: None for h in DEFAULT_HEADERS}]
    if args.dry_run:
        print(f"Would replace Posts body with {len(dates)} rows for {args.month}.")
        return
    backup = Path(config["workbook_path"]).with_suffix(f".backup-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx")
    shutil.copy2(config["workbook_path"], backup)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for index, day in enumerate(dates, start=2):
        template = existing_by_date.get(day)
        if not template:
            prior_cycle_day = day - dt.timedelta(days=len(dates))
            template = existing_by_date.get(prior_cycle_day)
        if not template:
            template = templates[(index - 2) % len(templates)]
        image = curated_image_for_date(config, day)
        for header in DEFAULT_HEADERS:
            value = template.get(header)
            if header == "Date":
                value = dt.datetime.combine(day, dt.time())
            elif header == "AssetSource" and image:
                value = "Curated OneDrive"
            elif header == "AssetIdOrDescription" and image:
                value = str(image)
            elif header == "Status":
                value = "Ready" if image else config.get("prepared_status", "Needs approval")
            elif header == "CaptionDraft":
                value = None
            elif header == "ImageLink":
                value = None
            elif header == "Posted":
                value = False
            elif header in {"PostedAt", "GBPPostUrl", "Notes"}:
                value = None
            ws.cell(index, cols[header]).value = value
    wb.save(config["workbook_path"])
    log(config, f"Created {args.month} schedule with {len(dates)} rows. Backup: {backup}")


def supabase_public_url(config: dict[str, Any], path: str) -> str:
    base = config["supabase_url"].rstrip("/")
    bucket = urllib.parse.quote(config["supabase_bucket"].strip("/"))
    return f"{base}/storage/v1/object/public/{bucket}/{urllib.parse.quote(path)}"


def upload_supabase(path: Path, config: dict[str, Any]) -> str:
    url = config.get("supabase_url", "").rstrip("/")
    key = config.get("supabase_service_role_key", "")
    bucket = config.get("supabase_bucket", "")
    if not url or not key or not bucket:
        raise RuntimeError("Supabase URL, service role key, and bucket are required.")
    object_name = f"{dt.date.today().isoformat()}/{path.name}"
    endpoint = f"{url}/storage/v1/object/{urllib.parse.quote(bucket)}/{urllib.parse.quote(object_name)}"
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    request = urllib.request.Request(
        endpoint,
        data=path.read_bytes(),
        headers={
            "Authorization": f"Bearer {key}",
            "apikey": key,
            "Content-Type": mime,
            "x-upsert": "true",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=120):
        pass
    return supabase_public_url(config, object_name)


def google_access_token(config: dict[str, Any]) -> str:
    refresh = config.get("google_refresh_token")
    client_id = config.get("google_client_id")
    client_secret = config.get("google_client_secret")
    if not refresh or not client_id or not client_secret:
        raise RuntimeError("Google OAuth client id, secret, and refresh token are required.")
    payload = urllib.parse.urlencode(
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "refresh_token": refresh,
            "grant_type": "refresh_token",
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        "https://oauth2.googleapis.com/token",
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))
    return data["access_token"]


def update_local_config(updates: dict[str, Any]) -> None:
    config_path = CONFIG_LOCAL
    if not config_path.exists():
        shutil.copy2(CONFIG_EXAMPLE, config_path)
    with config_path.open("r", encoding="utf-8") as handle:
        config = json.load(handle)
    config.update({key: value for key, value in updates.items() if value is not None})
    save_json(config_path, config)


def google_auth_url(args: argparse.Namespace) -> None:
    config = load_config()
    client_id = args.client_id or config.get("google_client_id")
    if not client_id:
        raise SystemExit("Set google_client_id in config.local.json or pass --client-id.")
    if args.client_id:
        update_local_config({"google_client_id": args.client_id, "google_client_secret": args.client_secret})
    params = urllib.parse.urlencode(
        {
            "client_id": client_id,
            "redirect_uri": args.redirect_uri,
            "response_type": "code",
            "scope": "https://www.googleapis.com/auth/business.manage",
            "access_type": "offline",
            "prompt": "consent",
        }
    )
    print("Open this URL, approve access, then run google-auth-finish with the returned code:")
    print(f"https://accounts.google.com/o/oauth2/v2/auth?{params}")


def google_token_from_code(config: dict[str, Any], code: str, redirect_uri: str) -> dict[str, Any]:
    client_id = config.get("google_client_id")
    client_secret = config.get("google_client_secret")
    if not client_id or not client_secret:
        raise RuntimeError("google_client_id and google_client_secret are required.")
    payload = urllib.parse.urlencode(
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "code": code,
            "grant_type": "authorization_code",
            "redirect_uri": redirect_uri,
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        "https://oauth2.googleapis.com/token",
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def google_get_json(url: str, token: str) -> dict[str, Any]:
    request = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def normalize_google_id(name_or_id: str) -> str:
    return name_or_id.rsplit("/", 1)[-1]


def google_auth_finish(args: argparse.Namespace) -> None:
    config = load_config()
    tokens = google_token_from_code(config, args.code, args.redirect_uri)
    refresh = tokens.get("refresh_token")
    access = tokens.get("access_token")
    if not access:
        raise SystemExit("Google did not return an access token.")
    accounts = google_get_json("https://mybusinessaccountmanagement.googleapis.com/v1/accounts", access).get("accounts", [])
    print("Google Business Profile accounts:")
    for item in accounts:
        print(f"- {item.get('name')} {item.get('accountName') or item.get('type') or ''}")
    account_id = args.account_id
    if not account_id and accounts:
        account_id = normalize_google_id(str(accounts[0].get("name", "")))
    location_id = args.location_id
    if account_id and not location_id:
        url = (
            "https://mybusinessbusinessinformation.googleapis.com/v1/"
            f"accounts/{account_id}/locations?readMask=name,title,storefrontAddress"
        )
        try:
            locations = google_get_json(url, access).get("locations", [])
        except urllib.error.HTTPError as exc:
            locations = []
            print(f"Could not list locations yet: HTTP {exc.code}")
        print("Google Business Profile locations:")
        for item in locations:
            print(f"- {item.get('name')} {item.get('title') or ''}")
        if locations:
            location_id = normalize_google_id(str(locations[0].get("name", "")))
    update_local_config(
        {
            "google_refresh_token": refresh or config.get("google_refresh_token"),
            "google_account_id": normalize_google_id(account_id) if account_id else config.get("google_account_id"),
            "google_location_id": normalize_google_id(location_id) if location_id else config.get("google_location_id"),
        }
    )
    print(f"Saved Google auth details to {CONFIG_LOCAL}")


# ---------------------------------------------------------------------------
# GBP post content-policy gate.
#
# Ported from scripts/gbp-poster/policy-check.mjs in barnscarter-ops/SEO-Agents-App.
#
# Incident driving this gate: 2026-07-31 -- three scheduled posts carried
# "(469) 863-9804" in the caption; Google rejected all three within minutes and
# disabled posting profile-wide. This gate makes that class of failure
# impossible to ship again: post_workbook_row() fails closed on any violation.
#
# Grounded in Google's published policies (reviewed 2026-08-01):
#   - Posts content policy: "We do not allow your post content to include a
#     phone number." Links/CTAs belong in the post's button, not the text.
#   - Prohibited & restricted content: offensive/sexual content, regulated
#     goods, personal info, off-topic, repetitive/duplicate content.
# ---------------------------------------------------------------------------

CAPTION_MAX = 1500

# Phone patterns: (469) 863-9804 / 469-863-9804 / 469.863.9804 / 4698639804 /
# +1 469 863 9804 / 863-9804 / vanity 555-555-TREE.
_PHONE_PATTERNS = [
    re.compile(r"(\+?1[\s.\-]?)?(\(\d{3}\)|\b\d{3})[\s.\-]?\d{3}[\s.\-]?\d{4}\b"),
    re.compile(r"\b\d{3}[\s.\-]\d{4}\b"),
    re.compile(r"\b\d{3}[\s.\-]\d{3}[\s.\-][A-Za-z]{4}\b"),
]
_URL_PATTERN = re.compile(
    r"(https?://|www\.)\S+|\b[a-z0-9][a-z0-9-]*\.(com|net|org|io|co|us|biz|info|me|app|site|online)\b",
    re.IGNORECASE,
)
_EMAIL_PATTERN = re.compile(r"\b[\w.+-]+@[\w-]+\.[a-z]{2,}\b", re.IGNORECASE)
# Street address or city/state+zip in text (addresses belong on the profile).
_ADDRESS_PATTERNS = [
    re.compile(
        r"\b\d{1,6}\s+(?:[A-Z][A-Za-z]+\s+){1,3}"
        r"(St|Street|Ave|Avenue|Blvd|Boulevard|Rd|Road|Dr|Drive|Ln|Lane|Ct|Court|"
        r"Pkwy|Parkway|Hwy|Highway)\.?\b"
    ),
    re.compile(r"\b(Suite|Ste\.?|Unit|Apt\.?)\s*#?\s*\d+\b", re.IGNORECASE),
    re.compile(r"\b[A-Z]{2}\s+\d{5}(-\d{4})?\b"),
]
# "sex" is a documented auto-reject even in innocent phrases; the rest are
# plain profanity Google's filter treats as offensive content.
_BANNED_WORDS = re.compile(
    r"\b(sex\w*|fuck\w*|shit\w*|bitch\w*|damn(ed|it)?|porn\w*|nude|nudity)\b",
    re.IGNORECASE,
)
# Regulated goods/services (restricted content policy). An electrical
# contractor's posts should never mention these; any hit is a red flag.
_REGULATED = re.compile(
    r"\b(alcohol|beer|wine|liquor|vodka|whiskey|tobacco|cigarettes?|vapes?|vaping|"
    r"cannabis|marijuana|cbd|kratom|guns?|firearms?|ammo|ammunition|casinos?|"
    r"gambling|lottery|betting|pharmacy|pharmaceutical|prescription|opioids?)\b",
    re.IGNORECASE,
)


@dataclasses.dataclass
class PolicyViolation:
    rule: str
    detail: str


def policy_check_caption(
    caption: str, other_captions: list[str] | None = None
) -> list[PolicyViolation]:
    """Validate a post caption against Google's post content policy.

    Returns a list of violations (empty = compliant). Ported from
    checkPostPolicy() in policy-check.mjs; same rules, same order.
    """
    violations: list[PolicyViolation] = []
    caption = str(caption or "")

    if not caption.strip():
        violations.append(PolicyViolation("empty-caption", "Caption is empty."))
        return violations
    if len(caption) > CAPTION_MAX:
        violations.append(
            PolicyViolation(
                "caption-too-long",
                f"Caption is {len(caption)} chars (GBP limit {CAPTION_MAX}).",
            )
        )

    for pattern in _PHONE_PATTERNS:
        match = pattern.search(caption)
        if match:
            violations.append(
                PolicyViolation(
                    "phone-number",
                    f'Phone number in post text ("{match.group(0).strip()}") \u2014 '
                    "Google auto-rejects this. Use the CTA button; the profile's "
                    'verified number powers "Call now".',
                )
            )
            break

    match = _URL_PATTERN.search(caption)
    if match:
        violations.append(
            PolicyViolation(
                "url-in-text",
                f'URL/domain in post text ("{match.group(0).strip()}") \u2014 '
                "links belong in the CTA button, not the caption.",
            )
        )
    match = _EMAIL_PATTERN.search(caption)
    if match:
        violations.append(
            PolicyViolation(
                "email-in-text",
                f'Email address in post text ("{match.group(0)}").',
            )
        )
    for pattern in _ADDRESS_PATTERNS:
        match = pattern.search(caption)
        if match:
            violations.append(
                PolicyViolation(
                    "address-in-text",
                    f'Street address/zip in post text ("{match.group(0).strip()}") \u2014 '
                    "addresses belong on the profile, not in posts.",
                )
            )
            break

    match = _BANNED_WORDS.search(caption)
    if match:
        violations.append(
            PolicyViolation(
                "banned-word",
                f'"{match.group(0)}" triggers Google\u2019s offensive-content filter '
                "(even in innocent phrases).",
            )
        )
    match = _REGULATED.search(caption)
    if match:
        violations.append(
            PolicyViolation(
                "regulated-content",
                f'"{match.group(0)}" falls under Google\u2019s restricted '
                "products/services policy.",
            )
        )

    letters = re.sub(r"[^A-Za-z]", "", caption)
    if len(letters) > 30:
        upper = len(re.sub(r"[^A-Z]", "", caption))
        if upper / len(letters) > 0.6:
            violations.append(
                PolicyViolation(
                    "all-caps",
                    "Caption is mostly uppercase \u2014 reads as spam/shouting to the filter.",
                )
            )

    if other_captions:
        def _norm(value: str) -> str:
            return re.sub(r"\s+", " ", str(value or "")).strip().lower()

        self_norm = _norm(caption)
        if self_norm and any(_norm(other) == self_norm for other in other_captions):
            violations.append(
                PolicyViolation(
                    "duplicate-caption",
                    "Caption is identical to another post in the schedule \u2014 "
                    "Google rejects repetitive content.",
                )
            )

    return violations


def format_policy_violations(violations: list[PolicyViolation]) -> str:
    return "\n".join(f"  [{v.rule}] {v.detail}" for v in violations)


def create_google_post(caption: str, image_url: str, config: dict[str, Any]) -> str:
    account = config.get("google_account_id")
    location = config.get("google_location_id")
    if not account or not location:
        raise RuntimeError("Google account/location IDs are required.")
    token = google_access_token(config)
    endpoint = f"https://mybusiness.googleapis.com/v4/accounts/{account}/locations/{location}/localPosts"
    payload = {
        "languageCode": "en-US",
        "summary": caption,
        "topicType": "STANDARD",
        "media": [{"mediaFormat": "PHOTO", "sourceUrl": image_url}],
    }
    request = urllib.request.Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=90) as response:
        data = json.loads(response.read().decode("utf-8"))
    return data.get("searchUrl") or data.get("name") or ""


def post_workbook_row(
    config: dict[str, Any],
    wb: Any,
    ws: Any,
    cols: dict[str, int],
    row: int,
    target: dt.date,
    dry_run: bool,
) -> bool:
    status = str(ws.cell(row, cols["Status"]).value or "")
    posted = bool(ws.cell(row, cols["Posted"]).value)
    if status != config.get("approval_status", "Approved") or posted:
        print(f"Row {row} is not postable. Status={status!r} Posted={posted}")
        return False
    caption = str(ws.cell(row, cols["CaptionDraft"]).value or "").strip()
    asset = str(ws.cell(row, cols["AssetIdOrDescription"]).value or "").strip()
    image_link = str(ws.cell(row, cols["ImageLink"]).value or "").strip()
    if not caption:
        print(f"Row {row} skipped: CaptionDraft is required before posting.")
        return False
    violations = policy_check_caption(caption)
    if violations:
        detail = format_policy_violations(violations)
        print(f"Row {row} BLOCKED by content-policy gate:\n{detail}")
        if dry_run:
            return False
        log(config, f"Policy gate blocked GBP row {row} for {target.isoformat()}:\n{detail}")
        ws.cell(row, cols["Notes"]).value = (
            f"BLOCKED by content-policy gate at {now_stamp()}:\n{detail}"
        )
        wb.save(config["workbook_path"])
        return False
    if not image_link:
        image_path = Path(asset)
        if not image_path.exists():
            image_path = curated_image_for_date(config, target) or image_path
        if dry_run:
            print(f"Would upload image: {image_path}")
            image_link = "https://example.invalid/dry-run-image.jpg"
        else:
            image_link = upload_supabase(image_path, config)
            ws.cell(row, cols["ImageLink"]).value = image_link
            wb.save(config["workbook_path"])
    if dry_run:
        print(f"Would create GBP post for row {row}: {caption[:100]}... image={image_link}")
        return True
    post_url = create_google_post(caption, image_link, config)
    ws.cell(row, cols["Posted"]).value = True
    ws.cell(row, cols["PostedAt"]).value = dt.datetime.now()
    ws.cell(row, cols["GBPPostUrl"]).value = post_url
    ws.cell(row, cols["Notes"]).value = f"Posted by grizzly-gbp-poster at {now_stamp()}"
    wb.save(config["workbook_path"])
    log(config, f"Posted GBP row {row} for {target.isoformat()}: {post_url}")
    return True


def post_approved(args: argparse.Namespace) -> None:
    config = load_config()
    target = dt.date.fromisoformat(args.date) if args.date else dt.date.today()
    wb, ws, cols = open_posts_sheet(config)
    if getattr(args, "all_due", False):
        count = 0
        for row_index in range(2, ws.max_row + 1):
            date_value = row_date(ws.cell(row_index, cols["Date"]).value)
            if date_value and date_value <= target:
                count += int(post_workbook_row(config, wb, ws, cols, row_index, date_value, args.dry_run))
        print(f"Approved due rows processed: {count}")
        return
    row = find_row_by_date(ws, cols, target)
    if not row:
        raise SystemExit(f"No Posts row found for {target.isoformat()}.")
    post_workbook_row(config, wb, ws, cols, row, target, args.dry_run)


def safe_archive_name(date_value: str, topic: str, original: Path) -> str:
    return f"{date_value}-{slugify(topic)}-{short_hash(original)}{original.suffix.lower()}"


def set_online_only(path: Path) -> None:
    subprocess.run(["attrib", "+U", "-P", str(path)], check=False, capture_output=True, text=True)


def archive_month(args: argparse.Namespace) -> None:
    config = load_config()
    wb, ws, cols = open_posts_sheet(config)
    archive_root = Path(config["archive_folder"]) / args.month
    rows: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        date_value = row_date(ws.cell(row, cols["Date"]).value)
        if not date_value or date_value.strftime("%Y-%m") != args.month:
            continue
        posted = bool(ws.cell(row, cols["Posted"]).value)
        gbp_url = str(ws.cell(row, cols["GBPPostUrl"]).value or "").strip()
        if not posted and not gbp_url:
            continue
        asset = str(ws.cell(row, cols["AssetIdOrDescription"]).value or "").strip()
        image_link = str(ws.cell(row, cols["ImageLink"]).value or "").strip()
        source = Path(asset)
        if not source.exists():
            fallback = curated_image_for_date(config, date_value)
            source = fallback if fallback else source
        if not source.exists():
            log(config, f"ARCHIVE missing source for row {row}: {asset}")
            continue
        target_name = safe_archive_name(date_value.isoformat(), str(ws.cell(row, cols["Topic"]).value or "gbp-post"), source)
        target = archive_root / target_name
        metadata = {
            "posted_date": date_value.isoformat(),
            "topic": ws.cell(row, cols["Topic"]).value,
            "cta": ws.cell(row, cols["CTA"]).value,
            "source_file": str(source),
            "archived_file": str(target),
            "image_link": image_link,
            "gbp_post_url": gbp_url,
            "usage_status": "posted",
            "archived_at": now_stamp(),
        }
        rows.append(metadata)
        if args.dry_run:
            print(f"Would archive {source} -> {target}")
            continue
        archive_root.mkdir(parents=True, exist_ok=True)
        if target.exists():
            target = archive_root / f"{target.stem}-{int(time.time())}{target.suffix}"
            metadata["archived_file"] = str(target)
        shutil.move(str(source), str(target))
        sidecar = target.with_suffix(target.suffix + ".json")
        save_json(sidecar, metadata)
        if args.online_only:
            set_online_only(target)
            set_online_only(sidecar)
    if not args.dry_run and rows:
        archive_root.mkdir(parents=True, exist_ok=True)
        csv_path = archive_root / "usage-log.csv"
        exists = csv_path.exists()
        with csv_path.open("a", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            if not exists:
                writer.writeheader()
            writer.writerows(rows)
        if args.online_only:
            set_online_only(csv_path)
    print(f"Archive rows processed: {len(rows)}")


def run_daily(args: argparse.Namespace) -> None:
    scan_args = argparse.Namespace(dry_run=False, recent_days=None)
    scan_photos(scan_args)
    prepare_args = argparse.Namespace(date=args.date, dry_run=False)
    prepare_today(prepare_args)
    post_args = argparse.Namespace(date=args.date, dry_run=False, all_due=True)
    post_approved(post_args)


def main() -> None:
    parser = argparse.ArgumentParser(description="Grizzly GBP posting workflow.")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("init-config").set_defaults(func=init_config)

    scan = sub.add_parser("scan-photos")
    scan.add_argument("--dry-run", action="store_true")
    scan.add_argument("--recent-days", type=int)
    scan.set_defaults(func=scan_photos)

    prep = sub.add_parser("prepare-today")
    prep.add_argument("--date")
    prep.add_argument("--dry-run", action="store_true")
    prep.set_defaults(func=prepare_today)

    post = sub.add_parser("post-approved")
    post.add_argument("--date")
    post.add_argument("--dry-run", action="store_true")
    post.add_argument("--all-due", action="store_true", help="Post all approved, unposted rows dated on or before --date/today.")
    post.set_defaults(func=post_approved)

    sched = sub.add_parser("create-month-schedule")
    sched.add_argument("--month", required=True, help="YYYY-MM")
    sched.add_argument("--days", type=int)
    sched.add_argument("--dry-run", action="store_true")
    sched.set_defaults(func=create_month_schedule)

    arch = sub.add_parser("archive-month")
    arch.add_argument("--month", required=True, help="YYYY-MM")
    arch.add_argument("--dry-run", action="store_true")
    arch.add_argument("--online-only", action="store_true")
    arch.set_defaults(func=archive_month)

    daily = sub.add_parser("run-daily")
    daily.add_argument("--date")
    daily.set_defaults(func=run_daily)

    auth_url = sub.add_parser("google-auth-url")
    auth_url.add_argument("--client-id")
    auth_url.add_argument("--client-secret")
    auth_url.add_argument("--redirect-uri", default="http://localhost")
    auth_url.set_defaults(func=google_auth_url)

    auth_finish = sub.add_parser("google-auth-finish")
    auth_finish.add_argument("--code", required=True)
    auth_finish.add_argument("--redirect-uri", default="http://localhost")
    auth_finish.add_argument("--account-id")
    auth_finish.add_argument("--location-id")
    auth_finish.set_defaults(func=google_auth_finish)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    try:
        main()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise SystemExit(f"HTTP error {exc.code}: {body}") from exc
